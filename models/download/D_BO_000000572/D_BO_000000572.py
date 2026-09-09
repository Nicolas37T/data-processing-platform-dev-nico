import os
import re
import json
import urllib.request
from datetime import datetime
from openpyxl import load_workbook

from models.download.Download_Base import Download_Base
from models.download.tools.download_tools import format_date


class D_BO_000000572(Download_Base):
    """Robot class to extract URLs and validate 'Inyecciones STI' reports."""

    def get_file_url(
        self,
        main_url: str,
        updated_to: str,
        key_words: str = "",
        format: str = "%Y-%m-%d",
    ) -> list[str]:
     
        urls: list[str] = []

        try:
            ref_date = datetime.strptime(updated_to, format)
        except ValueError as err:
            print(f"Error parsing updated_to date '{updated_to}': {err}")
            return []

        try:
            # Category ID 258 corresponds to 'Inyecciones STI'
            api_endpoint = (
                "https://www.cndc.bo/wp-json/cndc/v1/estadisticas/documentos"
                "?categoria_id=258&agrupado=true"
            )

            req = urllib.request.Request(
                api_endpoint,
                headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"},
            )

            with urllib.request.urlopen(req, timeout=30) as resp:
                data = json.loads(resp.read().decode("utf-8"))

            grupos = data.get("grupos", [])
            for grupo in grupos:
                docs = grupo.get("docs", [])
                for doc in docs:
                    archivo_url = doc.get("archivo_url")
                    fecha_str = doc.get("fecha") or doc.get("periodo")  # YYYY-MM-DD
                    titulo = doc.get("titulo", "")

                    if not archivo_url:
                        continue

                    if key_words and key_words.lower() not in titulo.lower():
                        continue

                    if fecha_str:
                        try:
                            doc_date = datetime.strptime(fecha_str, "%Y-%m-%d")
                            if doc_date > ref_date:
                                urls.append(archivo_url)
                        except ValueError:
                            urls.append(archivo_url)
                    else:
                        urls.append(archivo_url)

        except Exception as exc:
            print(f"Error in get_file_url: {exc}")
            return []

        return urls

    def compare_files(
        self,
        files_paths: list[dict],
        updated_to: str,
        format: str = "%Y-%m-%d",
    ):
 
        valid_files: list[dict] = []

        try:
            ref_date = datetime.strptime(updated_to, format)
        except ValueError as err:
            print(f"Error parsing updated_to date '{updated_to}': {err}")
            return False

        months_es = {
            "enero": 1,
            "febrero": 2,
            "marzo": 3,
            "abril": 4,
            "mayo": 5,
            "junio": 6,
            "julio": 7,
            "agosto": 8,
            "septiembre": 9,
            "octubre": 10,
            "noviembre": 11,
            "diciembre": 12,
        }

        for entry in files_paths:
            tmp_path = entry.get("tmp_path", "")
            download_url = entry.get("download_url", "")

            if not os.path.isfile(tmp_path):
                print(f"File not found: {tmp_path}")
                continue

            try:
                file_date = None

                # Extract date from filename pattern inyecsti_DDMMYY.xlsx
                basename = os.path.basename(tmp_path)
                fname_match = re.search(r"(\d{6})\.\w+$", basename)

                if fname_match:
                    date_str = fname_match.group(1)
                    day = int(date_str[0:2])
                    month = int(date_str[2:4])
                    year = 2000 + int(date_str[4:6])
                    file_date = datetime(year, month, day)

                # Fallback: Read Excel content if date not found in filename
                if file_date is None:
                    wb = load_workbook(tmp_path, read_only=True, data_only=True)
                    ws = wb.active

                    for row in ws.iter_rows(
                        min_row=1, max_row=10, max_col=10, values_only=True
                    ):
                        for cell_value in row:
                            if cell_value is None:
                                continue

                            if isinstance(cell_value, datetime):
                                file_date = cell_value
                                break

                            cell_str = str(cell_value)
                            m = re.search(
                                r"(\d{1,2})\s+de\s+(\w+)\s+de\s+(\d{4})",
                                cell_str,
                                re.IGNORECASE,
                            )
                            if m:
                                d = int(m.group(1))
                                mn = months_es.get(m.group(2).lower())
                                y = int(m.group(3))
                                if mn:
                                    file_date = datetime(y, mn, d)
                                    break

                            m2 = re.search(r"(\d{2})/(\d{2})/(\d{4})", cell_str)
                            if m2:
                                file_date = datetime(
                                    int(m2.group(3)),
                                    int(m2.group(2)),
                                    int(m2.group(1)),
                                )
                                break

                        if file_date:
                            break

                    wb.close()

                if file_date is None:
                    print(f"Could not extract date from file: {tmp_path}")
                    continue

                if file_date > ref_date:
                    result = {
                        "tmp_path": tmp_path,
                        "download_url": download_url,
                        "updated_to": file_date.strftime("%Y-%m-%d"),
                    }
                    valid_files.append(result)

            except Exception as file_err:
                print(f"Error reading file '{tmp_path}': {file_err}")
                continue

        # If no valid files found, return False as required by platform
        if not valid_files:
            return False

        return valid_files