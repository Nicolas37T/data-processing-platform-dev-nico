"""Conversion robot for BCB exchange rates and UFV values."""

import os
import re
import traceback
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

import pandas as pd
from openpyxl import load_workbook
from unidecode import unidecode

from models.conversion.Conversion_Base import Conversion_Base
from models.conversion.tools.conversion_tools import (
    get_xlsx_report_dataframe,
    to_numeric_datax,
)


class D_BO_000000419_06(Conversion_Base):
    """Convert BCB exchange rate and UFV daily information."""

    def extraction(
        self,
        file_path,
        key_words,
        template_path="",
        page_number=1,
        format="%Y-%m-%d",
    ):
        """Extract the target section and its latest daily week."""
        del template_path

        try:
            extension = os.path.splitext(str(file_path))[1].lower()
            if extension != ".xlsx":
                raise ValueError(f"Unsupported file extension: {extension}")

            normalized_keywords = unidecode(
                str(key_words or "")
            ).lower()
            if normalized_keywords and not all(
                term in normalized_keywords
                for term in ("informacion", "estadistica")
            ):
                raise ValueError(
                    "The configured keywords do not match the report."
                )

            report = get_xlsx_report_dataframe(
                file_path=file_path,
                key_words="tipos cambio valor ufv",
                page_number=page_number,
            )
            if not report:
                raise ValueError(
                    "The exchange rate and UFV table was not found."
                )

            raw_dataframe, found_page_number = report
            workbook = load_workbook(file_path, data_only=True)
            worksheet = workbook.worksheets[found_page_number - 1]
            section_title = "Tipos de cambio y valor de la UFV"

            section_row = None
            section_column = None
            for row_index, row in raw_dataframe.iterrows():
                for column in raw_dataframe.columns:
                    value = row[column]
                    if pd.isna(value):
                        continue
                    normalized_value = unidecode(str(value)).lower()
                    if all(
                        term in normalized_value
                        for term in ("tipos de cambio", "valor", "ufv")
                    ):
                        section_row = row_index
                        section_column = int(column)
                        break
                if section_row is not None:
                    break

            if section_row is None or section_column is None:
                raise ValueError("The requested report section was not found.")

            variable_column = section_column + 1
            if variable_column not in raw_dataframe.columns:
                raise ValueError("The VARIABLES column was not found.")

            section_end_row = len(raw_dataframe)
            for row_index in raw_dataframe.index:
                if row_index <= section_row:
                    continue
                section_value = raw_dataframe.at[
                    row_index,
                    section_column,
                ]
                if not pd.isna(section_value):
                    section_end_row = row_index
                    break

            week_row = None
            week_start_column = None
            for row_index in raw_dataframe.index:
                if row_index >= section_row:
                    break
                for column in raw_dataframe.columns:
                    value = raw_dataframe.at[row_index, column]
                    if pd.isna(value):
                        continue
                    if re.fullmatch(
                        r"Semana\s+\d+",
                        str(value).strip(),
                        flags=re.IGNORECASE,
                    ):
                        if (
                            week_start_column is None
                            or int(column) > week_start_column
                        ):
                            week_row = row_index
                            week_start_column = int(column)

            if week_row is None or week_start_column is None:
                raise ValueError("The latest weekly group was not found.")

            date_row = None
            date_columns = []
            date_values = {}
            for row_index in raw_dataframe.index:
                if row_index <= week_row or row_index >= section_row:
                    continue

                candidate_dates = {}
                for column in raw_dataframe.columns:
                    if int(column) < week_start_column:
                        continue
                    value = raw_dataframe.at[row_index, column]
                    if pd.isna(value):
                        continue
                    if not isinstance(
                        value,
                        (str, date, datetime, pd.Timestamp),
                    ):
                        continue
                    parsed_date = pd.to_datetime(
                        value,
                        errors="coerce",
                    )
                    if not pd.isna(parsed_date):
                        candidate_dates[int(column)] = parsed_date.strftime(
                            format
                        )

                if len(candidate_dates) > len(date_columns):
                    date_row = row_index
                    date_columns = list(candidate_dates)
                    date_values = candidate_dates

            if date_row is None or not date_columns:
                raise ValueError(
                    "Daily columns for the latest week were not found."
                )

            source_rows = []
            previous_row = section_row - 1
            if previous_row >= 0:
                previous_label = raw_dataframe.at[
                    previous_row,
                    variable_column,
                ]
                normalized_previous_label = unidecode(
                    str(previous_label or "")
                ).lower()
                if all(
                    term in normalized_previous_label
                    for term in (
                        "transferencias del exterior",
                        "sistema financiero",
                    )
                ):
                    source_rows.append(previous_row)

            source_rows.extend(
                range(section_row + 1, section_end_row)
            )
            records = []
            for row_index in source_rows:
                raw_label = raw_dataframe.at[row_index, variable_column]
                if pd.isna(raw_label):
                    continue

                label = re.sub(r"\s+", " ", str(raw_label)).strip()
                label = re.sub(
                    r"trav.s",
                    "través",
                    label,
                    flags=re.IGNORECASE,
                )
                label = re.sub(
                    r"Bols.n",
                    "Bolsín",
                    label,
                    flags=re.IGNORECASE,
                )
                label = re.sub(
                    r"d.lar",
                    "dólar",
                    label,
                    flags=re.IGNORECASE,
                )
                label = re.sub(
                    r"^.ndice",
                    "Índice",
                    label,
                    flags=re.IGNORECASE,
                )
                label = re.sub(
                    r"d.a",
                    "día",
                    label,
                    flags=re.IGNORECASE,
                )
                label = re.sub(
                    r"h.bil",
                    "hábil",
                    label,
                    flags=re.IGNORECASE,
                )
                label = re.sub(
                    r".ltimo",
                    "último",
                    label,
                    flags=re.IGNORECASE,
                )

                for column in date_columns:
                    source_value = raw_dataframe.at[row_index, column]
                    normalized_value = pd.NA
                    if not pd.isna(source_value):
                        source_cell = worksheet.cell(
                            row=row_index + 1,
                            column=column + 1,
                        )
                        number_format = (
                            source_cell.number_format or "General"
                        )
                        format_sections = number_format.split(";")
                        zero_as_dash = (
                            len(format_sections) >= 3
                            and '"-"' in format_sections[2]
                        )

                        if str(source_value).strip() == "-":
                            normalized_value = "-"
                        else:
                            if isinstance(source_value, (int, float)):
                                numeric_value = float(source_value)
                            else:
                                converted_value = to_numeric_datax(
                                    pd.Series([source_value]),
                                    ",",
                                ).iloc[0]
                                if pd.isna(converted_value):
                                    raise ValueError(
                                        f"Invalid numeric value for {label}."
                                    )
                                numeric_value = float(converted_value)

                            if numeric_value == 0 and zero_as_dash:
                                normalized_value = "-"
                            else:
                                positive_format = format_sections[0]
                                decimal_match = re.search(
                                    r"\.([0#]+)",
                                    positive_format,
                                )
                                decimal_places = (
                                    len(decimal_match.group(1))
                                    if decimal_match
                                    else 0
                                )
                                quantizer = Decimal("1").scaleb(
                                    -decimal_places
                                )
                                rounded_value = Decimal(
                                    str(numeric_value)
                                ).quantize(
                                    quantizer,
                                    rounding=ROUND_HALF_UP,
                                )
                                if decimal_places == 0:
                                    normalized_value = int(rounded_value)
                                else:
                                    normalized_value = float(rounded_value)

                    records.append(
                        {
                            "nv1": section_title,
                            "nv2": label,
                            "fecha": date_values[column],
                            "valor": normalized_value,
                        }
                    )

            if not records:
                raise ValueError("No target table records were extracted.")

            dataframe = pd.DataFrame.from_records(
                records,
                columns=["nv1", "nv2", "fecha", "valor"],
            )
            if dataframe.duplicated(
                subset=["nv1", "nv2", "fecha"]
            ).any():
                raise ValueError("Duplicate exchange rate records were found.")

            # Drop empty / null values
            dataframe = dataframe.dropna(subset=["valor"])

            metadata = {
                "file_name": os.path.basename(file_path),
                "titles": [section_title],
                "page_number": int(found_page_number) if found_page_number and int(found_page_number) > 0 else 1,
            }
            return metadata, dataframe

        except Exception as error:
            print(f"Could not extract report from {file_path}: {error}")
            traceback.print_exc()
            return ""

    def validate_data_results(
        self,
        dataframe,
        decimal_separator,
        TOLERANCE=6.0,
    ):
        """Validate BCB exchange rate and UFV records."""
        del TOLERANCE

        try:
            required_columns = {"nv1", "nv2", "fecha", "valor"}
            if not isinstance(dataframe, pd.DataFrame) or dataframe.empty:
                return True
            if not required_columns.issubset(dataframe.columns):
                return True
            if dataframe[["nv1", "nv2", "fecha"]].isna().any().any():
                return True
            if dataframe.duplicated(
                subset=["nv1", "nv2", "fecha"]
            ).any():
                return True
            if pd.to_datetime(
                dataframe["fecha"],
                format="%Y-%m-%d",
                errors="coerce",
            ).isna().any():
                return True

            source_values = dataframe["valor"].dropna()
            dash_mask = source_values.astype(str).str.strip().eq("-")
            values_to_check = source_values.loc[~dash_mask]
            if values_to_check.empty:
                return True

            numeric_values = pd.to_numeric(
                values_to_check,
                errors="coerce",
            )
            if numeric_values.isna().any():
                numeric_values = to_numeric_datax(
                    values_to_check,
                    decimal_separator,
                )
            return bool(numeric_values.isna().any())

        except Exception as error:
            print(f"Could not validate exchange rate report: {error}")
            traceback.print_exc()
            return True


# Compatibility alias
Robot = D_BO_000000419_06
