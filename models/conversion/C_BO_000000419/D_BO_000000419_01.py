"""Conversion robot for BCB external operations report D_BO_000000419_01."""

import os
import re
import traceback
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

import pandas as pd
from openpyxl import load_workbook
from unidecode import unidecode

from models.conversion.Conversion_Base import Conversion_Base
from models.conversion.tools.conversion_tools import to_numeric_datax


class D_BO_000000419_01(Conversion_Base):
    """Convert BCB external operations report."""

    def extraction(
        self,
        file_path,
        key_words,
        template_path="",
        page_number=1,
        format="%Y-%m-%d",
    ):
        """Extract external operations table and its latest daily week."""
        del template_path

        def normalize_superscripts(text):
            if not text:
                return ""
            trans_table = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")
            return text.translate(trans_table).strip()

        def clean_numeric(val):
            if val is None or pd.isna(val):
                return None
            if isinstance(val, (int, float)):
                return float(val)
            s = str(val).strip()
            if not s or s == "-":
                return None
            s = s.replace(" ", "").replace("%", "")
            if "," in s and "." in s:
                if s.find(",") < s.find("."):
                    s = s.replace(",", "")
                else:
                    s = s.replace(".", "").replace(",", ".")
            elif "," in s:
                if re.match(r"^-?\d+,\d{3}$", s):
                    s = s.replace(",", "")
                else:
                    s = s.replace(",", ".")
            try:
                return float(s)
            except ValueError:
                return None

        def determine_level(var_name, raw_indent):
            lower_name = unidecode(var_name).lower()
            if raw_indent > 0:
                return min(int(raw_indent), 3)
            if "d/c en portafolio" in lower_name:
                return 3
            if re.match(r"^[a-b]\.\s", lower_name) or re.match(r"^[d-g]\.\s", lower_name):
                return 2
            if lower_name in {"divisas", "deg", "posicion con el fmi"}:
                return 2
            if "oro en toneladas" in lower_name or re.match(r"^c\.\s", lower_name) or lower_name == "oro":
                return 1
            if "reservas de oro neto de pasivos (d + e - f)" in lower_name:
                return 1
            if "d/c recursos de alta liquidez" in lower_name or "d/c fondo ral" in lower_name:
                return 1
            return 0

        try:
            extension = os.path.splitext(str(file_path))[1].lower()
            if extension != ".xlsx":
                raise ValueError(f"Unsupported file extension: {extension}")

            normalized_keywords = unidecode(str(key_words or "")).lower()
            if normalized_keywords and not all(
                term in normalized_keywords
                for term in ("informacion", "estadistica")
            ):
                raise ValueError("The configured keywords do not match the report.")

            excel_file_obj = pd.ExcelFile(file_path)
            sheet_idx = (page_number - 1) if (0 <= page_number - 1 < len(excel_file_obj.sheet_names)) else 0
            found_page_number = sheet_idx + 1
            raw_dataframe = pd.read_excel(file_path, sheet_name=sheet_idx, header=None)

            workbook = load_workbook(file_path, data_only=True)
            worksheet = workbook.worksheets[sheet_idx]

            section_row = None
            section_column = None
            section_title = "Operaciones con el exterior (saldos y flujos en millones de USD)"
            document_title = "INFORMACIÓN ESTADÍSTICA SEMANAL (p)"

            for row_index, row in raw_dataframe.iterrows():
                for column in raw_dataframe.columns[:5]:
                    value = row[column]
                    if pd.isna(value):
                        continue
                    normalized_value = unidecode(str(value)).lower()
                    if "operaciones con el exterior" in normalized_value and "saldos" in normalized_value:
                        section_row = row_index
                        section_column = int(column)
                        section_title = str(value).strip()
                        break
                if section_row is not None:
                    break

            if section_row is None or section_column is None:
                raise ValueError("The requested report section was not found.")

            variable_column = section_column + 1
            if variable_column not in raw_dataframe.columns:
                raise ValueError("The VARIABLES column was not found.")

            section_end_row = len(raw_dataframe)
            for row_index in raw_dataframe.index:
                if row_index <= section_row:
                    continue
                section_value = raw_dataframe.at[row_index, section_column]
                if not pd.isna(section_value) and str(section_value).strip():
                    section_end_row = row_index
                    break

            # Find latest week column
            week_row = None
            week_start_column = None
            for row_index in raw_dataframe.index:
                if row_index >= section_row:
                    break
                for column in raw_dataframe.columns:
                    value = raw_dataframe.at[row_index, column]
                    if pd.isna(value):
                        continue
                    if re.fullmatch(r"Semana\s+\d+", str(value).strip(), flags=re.IGNORECASE):
                        if week_start_column is None or int(column) > week_start_column:
                            week_row = row_index
                            week_start_column = int(column)

            if week_row is None or week_start_column is None:
                raise ValueError("The latest weekly group was not found.")

            # Find date columns
            date_row = None
            date_columns = []
            date_values = {}
            for row_index in raw_dataframe.index:
                if row_index <= week_row or row_index >= section_row:
                    continue

                candidate_dates = {}
                for column in raw_dataframe.columns:
                    if int(column) < week_start_column:
                        continue
                    value = raw_dataframe.at[row_index, column]
                    if pd.isna(value):
                        continue
                    if not isinstance(value, (str, date, datetime, pd.Timestamp)):
                        continue
                    parsed_date = pd.to_datetime(value, errors="coerce")
                    if not pd.isna(parsed_date):
                        candidate_dates[int(column)] = parsed_date.strftime(format)

                if len(candidate_dates) > len(date_columns):
                    date_row = row_index
                    date_columns = list(candidate_dates)
                    date_values = candidate_dates

            if date_row is None or not date_columns:
                raise ValueError("Daily columns for the latest week were not found.")

            # Opción A: Si la última semana tiene menos de 5 días, buscar la fecha anterior
            if len(date_columns) < 5:
                first_daily_col = min(date_columns)
                prev_col = first_daily_col - 1
                found_prev_col = None
                found_prev_date = None

                while prev_col >= 0:
                    for r_check in (date_row, week_row):
                        val_candidate = raw_dataframe.at[r_check, prev_col]
                        if not pd.isna(val_candidate):
                            parsed_prev = pd.to_datetime(val_candidate, errors="coerce")
                            if not pd.isna(parsed_prev):
                                found_prev_col = prev_col
                                found_prev_date = parsed_prev.strftime(format)
                                break
                    if found_prev_col is not None:
                        break
                    prev_col -= 1

                if found_prev_col is not None and found_prev_col not in date_columns:
                    date_columns.insert(0, found_prev_col)
                    date_values[found_prev_col] = found_prev_date

            date_columns.sort()

            nv_stack = [section_title, "", "", "", ""]
            records = []

            for row_index in range(section_row + 1, section_end_row):
                raw_cell_val = raw_dataframe.at[row_index, variable_column]
                if pd.isna(raw_cell_val):
                    continue

                raw_val = str(raw_cell_val).strip()
                if not raw_val:
                    continue

                source_cell = worksheet.cell(row=row_index + 1, column=variable_column + 1)
                raw_indent = source_cell.alignment.indent if (source_cell.alignment and source_cell.alignment.indent) else 0
                leading_spaces = len(str(raw_cell_val)) - len(str(raw_cell_val).lstrip(" "))
                if raw_indent == 0 and leading_spaces >= 2:
                    raw_indent = leading_spaces // 2

                clean_var = normalize_superscripts(raw_val)
                level = determine_level(clean_var, raw_indent)

                if level == 0:
                    nv_stack[1] = clean_var
                    nv_stack[2] = ""
                    nv_stack[3] = ""
                    nv_stack[4] = ""
                elif level == 1:
                    nv_stack[2] = clean_var
                    nv_stack[3] = ""
                    nv_stack[4] = ""
                elif level == 2:
                    nv_stack[3] = clean_var
                    nv_stack[4] = ""
                elif level >= 3:
                    nv_stack[4] = clean_var

                for column in date_columns:
                    source_value = raw_dataframe.at[row_index, column]
                    val_cleaned = clean_numeric(source_value)
                    if val_cleaned is not None:
                        records.append({
                            "nv1": nv_stack[0],
                            "nv2": nv_stack[1],
                            "nv3": nv_stack[2],
                            "nv4": nv_stack[3],
                            "nv5": nv_stack[4],
                            "fecha": date_values[column],
                            "valor": val_cleaned,
                        })

            if not records:
                raise ValueError("No target table records were extracted.")

            dataframe = pd.DataFrame.from_records(
                records,
                columns=[
                    "nv1",
                    "nv2",
                    "nv3",
                    "nv4",
                    "nv5",
                    "fecha",
                    "valor",
                ],
            )

            dataframe["valor"] = to_numeric_datax(dataframe["valor"], ",")
            dataframe = dataframe.dropna(subset=["valor"])

            metadata = {
                "file_name": os.path.basename(file_path),
                "titles": [document_title],
                "page_number": int(found_page_number),
            }
            return metadata, dataframe

        except Exception as error:
            print(f"Could not extract report from {file_path}: {error}")
            traceback.print_exc()
            return ""

    def validate_data_results(
        self,
        dataframe,
        decimal_separator=",",
        TOLERANCE=6.0,
    ):
        """Validate numerical integrity of extracted data."""
        del decimal_separator, TOLERANCE

        try:
            if dataframe is None or (
                isinstance(dataframe, pd.DataFrame) and dataframe.empty
            ):
                return False

            required_columns = {"fecha", "valor", "nv1", "nv2"}
            if not required_columns.issubset(dataframe.columns):
                return True

            numeric_values = pd.to_numeric(dataframe["valor"], errors="coerce")
            if numeric_values.isna().any():
                return True

            return False

        except Exception as error:
            print(f"Could not validate report data: {error}")
            traceback.print_exc()
            return True


Robot = D_BO_000000419_01
