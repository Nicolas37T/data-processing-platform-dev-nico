from models.conversion.Conversion_Base import Conversion_Base
"""
Module for extracting and validating report D_BO_000000463_01.

Code: D_BO_000000463_01
Name: Reservas Internacionales Netas del Banco Central de Bolivia
Page: 1
Periodicity: Monthly
Decimal Separator: ,
Conversion Factor: 1
Keywords: reservas internacionales
"""

import os
import re
import traceback
from datetime import datetime
from typing import Any, Dict, List, Tuple, Union

import pandas as pd
from openpyxl import load_workbook

from models.conversion.tools.conversion_tools import to_numeric_datax


class D_BO_000000463_01(Conversion_Base):
    """Conversion robot for report D_BO_000000463_01.

    Extracts 'Reservas Internacionales Netas del Banco Central de Bolivia' table
    from BCB Excel files, mapping hierarchical levels [nv1, nv2, fecha, valor].
    """

    def extraction(
        self,
        file_path: str,
        key_words: str = "",
        template_path: str = "",
        page_number: int = 1,
        format: str = "%Y-%m-%d",
    ) -> Union[Tuple[Dict[str, Any], pd.DataFrame], Tuple[str, pd.DataFrame]]:
        """Extract and normalize data from 'Reservas Internacionales Netas' Excel file.

        Parameters
        ----------
        file_path : str
            Absolute path of the source Excel file (.xlsx / .xls).
        key_words : str, optional
            Keywords to locate sheet or table (default '').
        template_path : str, optional
            Path to template file (if applicable).
        page_number : int, optional
            1-indexed sheet page number (default 1).
        format : str, optional
            Target date format string (default '%Y-%m-%d').

        Returns
        -------
        tuple[dict, pd.DataFrame] | tuple[str, pd.DataFrame]
            Tuple of (metadata_dict, df_melted) or ('', empty DataFrame) on failure.
        """
        def parse_date(val: Any) -> Union[str, None]:
            """Extract date record directly from table cell and normalize to YYYY-MM-DD."""
            if val is None:
                return None

            # Excel datetime objects come parsed as Python datetime
            if isinstance(val, datetime):
                # Use last day of month as cut-off date
                import calendar
                last_day = calendar.monthrange(val.year, val.month)[1]
                return val.replace(day=last_day).strftime("%Y-%m-%d")

            s = str(val).strip()
            if not s or s in {"-", "--", "---", "NA", "N/A"}:
                return None

            # Remove footnote markers or status flags like (p), (e), *, etc.
            s_clean = re.sub(r"\(.*?\)", "", s).replace("*", "").strip()

            # Month-year pattern like Jan-22, May-25, Jun-25, etc.
            match_my = re.match(r"^([A-Za-z]{3})[-_/\s](\d{2}|\d{4})$", s_clean)
            if match_my:
                import calendar
                month_str, year_str = match_my.group(1), match_my.group(2)
                if len(year_str) == 2:
                    year_str = "20" + year_str
                try:
                    month_num = datetime.strptime(month_str, "%b").month
                    last_day = calendar.monthrange(int(year_str), month_num)[1]
                    return f"{year_str}-{month_num:02d}-{last_day:02d}"
                except ValueError:
                    pass

            # Standard ISO format YYYY-MM-DD
            match_iso = re.match(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$", s_clean)
            if match_iso:
                return s_clean.replace("/", "-")

            # YYYY-MM only (no day): use last day of month
            match_ym = re.match(r"^(\d{4})[-/](\d{1,2})$", s_clean)
            if match_ym:
                import calendar
                y, m = int(match_ym.group(1)), int(match_ym.group(2))
                last_day = calendar.monthrange(y, m)[1]
                return f"{y}-{m:02d}-{last_day:02d}"

            return None

        def clean_numeric_cell(val: Any) -> Union[float, None]:
            """Parse cell numeric value into float, matching the document's displayed precision."""
            if val is None:
                return None

            num: float = None
            if isinstance(val, (int, float)):
                num = float(val)
            else:
                s = str(val).strip()
                if not s or s in {"-", "--", "---", "NA", "N/A", "ND"}:
                    return None

                # Strip footnote / estimate markers like (e), (p), *, etc.
                s = re.sub(r"\(.*?\)", "", s).replace("*", "").strip()
                s = s.replace(" ", "").replace("%", "")
                if not s:
                    return None

                # Guard: long strings (>30 chars) with letters are clearly not numbers (e.g. footer text)
                if len(s) > 30 and re.search(r"[A-Za-zÁÉÍÓÚáéíóú]", s):
                    return None

                # Handle scientific notation if present (e.g. 1.947e-7)
                if re.search(r"[-+]?\d*\.?\d+[eE][-+]?\d+", s):
                    try:
                        num = float(s)
                    except ValueError:
                        import traceback; traceback.print_exc()
                        pass
                else:
                    if "," in s and "." in s:
                        if s.rfind(",") > s.rfind("."):
                            s = s.replace(".", "").replace(",", ".")
                        else:
                            s = s.replace(",", "")
                    elif "," in s:
                        if re.match(r"^-?\d+,\d{3}$", s):
                            s = s.replace(",", "")
                        else:
                            s = s.replace(",", ".")

                    try:
                        num = float(s)
                    except ValueError:
                        import traceback; traceback.print_exc()
                        return None

            if num is None:
                return None

            # Values like 1.947e-7 (0.0000001947) displayed as 0.0 in Excel
            if abs(num) < 1e-4:
                return 0.0

            return round(num, 1)

        try:
            if not os.path.isfile(file_path):
                print(f"File not found: {file_path}")
                return "", pd.DataFrame()

            wb = load_workbook(file_path, data_only=True)
            if not wb.sheetnames:
                print(f"No sheets found in {file_path}")
                return "", pd.DataFrame()

            ws = wb.worksheets[0]
            if page_number <= len(wb.worksheets):
                ws = wb.worksheets[page_number - 1]

            # 1. Locate nv1 header row
            nv1_row = -1
            for r in range(1, min(ws.max_row + 1, 25)):
                row_texts = [str(ws.cell(row=r, column=c).value or "").upper() for c in range(1, ws.max_column + 1)]
                if any("RESERVAS BRUTAS" in t for t in row_texts) or any("RESERVAS NETAS" in t for t in row_texts):
                    nv1_row = r
                    break

            if nv1_row == -1:
                print(f"Could not locate nv1 header row in {file_path}")
                return "", pd.DataFrame()

            # 2. Locate data start row and header rows for nv2
            data_start_row = -1
            for r in range(nv1_row + 1, min(nv1_row + 15, ws.max_row + 1)):
                c1_val = str(ws.cell(row=r, column=1).value or "").strip()
                c2_val = str(ws.cell(row=r, column=2).value or "").strip()
                if re.search(r"[A-Za-zÁÉÍÓÚáéíóú]+[-_/\s]+\d{2,4}", c1_val) or re.search(r"\d{4}[-/]\d{1,2}", c1_val):
                    data_start_row = r
                    break
                if re.search(r"[A-Za-zÁÉÍÓÚáéíóú]+[-_/\s]+\d{2,4}", c2_val) or re.search(r"\d{4}[-/]\d{1,2}", c2_val):
                    data_start_row = r
                    break

            if data_start_row == -1:
                for r in range(nv1_row + 2, min(nv1_row + 10, ws.max_row + 1)):
                    row_nums = [ws.cell(row=r, column=c).value for c in range(3, ws.max_column + 1)]
                    if any(isinstance(v, (int, float)) for v in row_nums):
                        data_start_row = r
                        break

            if data_start_row == -1:
                print(f"Could not locate data start row in {file_path}")
                return "", pd.DataFrame()

            header_rows = list(range(nv1_row + 1, data_start_row))

            # 3. Map columns to (nv1, nv2) taking literal header text
            col_mapping: Dict[int, Tuple[str, str]] = {}
            curr_nv1 = None
            for c in range(1, ws.max_column + 1):
                val_nv1 = ws.cell(row=nv1_row, column=c).value
                if val_nv1:
                    txt1 = re.sub(r"\s+", " ", str(val_nv1)).strip().upper()
                    if "RESERVAS BRUTAS" in txt1:
                        curr_nv1 = "RESERVAS BRUTAS"
                    elif "OBLIGACIONES" in txt1:
                        curr_nv1 = "OBLIGACIONES"
                    elif "RESERVAS NETAS" in txt1:
                        curr_nv1 = "RESERVAS NETAS"

                parts = []
                for hr in header_rows:
                    val_h = ws.cell(row=hr, column=c).value
                    if val_h is not None:
                        s_h = str(val_h).strip()
                        if s_h and s_h not in parts:
                            parts.append(s_h)

                if parts and curr_nv1:
                    raw_nv2 = " ".join(parts)
                    raw_nv2 = re.sub(r"\s+", " ", raw_nv2).strip()
                    if any(k in raw_nv2.lower() for k in ["saldos", "fin de", "fecha"]):
                        continue
                    col_mapping[c] = (curr_nv1, raw_nv2)

            if not col_mapping:
                print(f"No columns mapped in {file_path}")
                return "", pd.DataFrame()

            # 4. Extract rows
            rows_list: List[Dict[str, Any]] = []
            stop_keywords = ["fuente", "nota", "tipo de cambio", "(1)", "en millones"]
            empty_count = 0

            for r in range(data_start_row, ws.max_row + 1):
                raw_c1 = ws.cell(row=r, column=1).value
                raw_c2 = ws.cell(row=r, column=2).value

                txt_check = f"{raw_c1 or ''} {raw_c2 or ''}".lower()
                if any(txt_check.strip().startswith(sk) for sk in stop_keywords):
                    break

                row_date = parse_date(raw_c1)
                if not row_date:
                    row_date = parse_date(raw_c2)

                if not row_date:
                    empty_count += 1
                    if empty_count >= 10:
                        break
                    continue
                empty_count = 0

                for col_idx, (cat_nv1, cat_nv2) in col_mapping.items():
                    raw_val = ws.cell(row=r, column=col_idx).value
                    clean_val = clean_numeric_cell(raw_val)
                    if clean_val is not None:
                        rows_list.append({
                            "nv1": cat_nv1,
                            "nv2": cat_nv2,
                            "fecha": row_date,
                            "valor": clean_val,
                        })

            if not rows_list:
                print(f"No numeric rows extracted from {file_path}")
                return "", pd.DataFrame()

            df_melted = pd.DataFrame(rows_list)

            # 5. Normalize valor column vectorized
            df_melted["valor"] = to_numeric_datax(
                serie=df_melted["valor"], decimal_separator=","
            )
            df_melted["valor"] = df_melted["valor"].apply(
                lambda v: 0.0 if (pd.isna(v) or abs(v) < 1e-4) else round(float(v), 1)
            )
            df_melted = df_melted.dropna(subset=["valor"])

            # 6. Extract titles from sheet
            report_title = "RESERVAS INTERNACIONALES NETAS DEL BANCO CENTRAL DE BOLIVIA"
            subtitle = "(En millones de $us)"
            for r in range(1, nv1_row):
                for c in range(1, ws.max_column + 1):
                    v_str = str(ws.cell(row=r, column=c).value or "").strip()
                    if "RESERVAS INTERNACIONALES NETAS" in v_str.upper():
                        report_title = re.sub(r"\s+", " ", re.sub(r"\(\d+\)", "", v_str)).strip()
                    elif "MILLONES" in v_str.upper():
                        subtitle = re.sub(r"\s+", " ", v_str).strip()

            metadata = {
                "file_name": os.path.basename(file_path),
                "titles": [report_title, subtitle],
                "page_number": page_number,
            }

            return (metadata, df_melted)

        except Exception as error:
            print(f"Could not extract report from {file_path}: {error}")
            traceback.print_exc()
            return "", pd.DataFrame()

    def validate_data_results(
        self,
        dataframe: pd.DataFrame,
        decimal_separator: str = ",",
        TOLERANCE: float = 6.0,
    ) -> bool:
        """Reconcile arithmetic sums of components against totals.

        Parameters
        ----------
        dataframe : pd.DataFrame
            Normalized DataFrame returned by extraction().
        decimal_separator : str
            Decimal separator character (default ',').
        TOLERANCE : float, optional
            Maximum acceptable difference threshold (default 6.0).

        Returns
        -------
        bool
            False on success / valid, True on error / discrepancy.
        """
        try:
            if dataframe is None or not isinstance(dataframe, pd.DataFrame) or dataframe.empty:
                return True

            required_cols = ["nv1", "nv2", "fecha", "valor"]
            if not all(col in dataframe.columns for col in required_cols):
                return True

            df = dataframe.copy()
            df["valor"] = to_numeric_datax(df["valor"], decimal_separator=decimal_separator)
            if df["valor"].isna().any():
                return True

            # 1. Reconcile RESERVAS BRUTAS: components vs TOTAL
            brutas_df = df[df["nv1"] == "RESERVAS BRUTAS"]
            if not brutas_df.empty:
                totals = brutas_df[brutas_df["nv2"] == "TOTAL"].groupby("fecha")["valor"].first()
                components = brutas_df[brutas_df["nv2"] != "TOTAL"].groupby("fecha")["valor"].sum()
                reconciled_brutas = pd.concat([totals, components], axis=1, keys=["total", "sum_comp"]).dropna()
                if not reconciled_brutas.empty:
                    diff_brutas = (reconciled_brutas["total"] - reconciled_brutas["sum_comp"]).abs()
                    max_diff_brutas = diff_brutas.max()
                    if max_diff_brutas > TOLERANCE:
                        print(f"Validation error in RESERVAS BRUTAS: max diff {max_diff_brutas} exceeds {TOLERANCE}")
                        return True

            # 2. Reconcile RESERVAS NETAS: NETAS == BRUTAS TOTAL - OBLIGACIONES
            obligaciones_df = df[df["nv1"] == "OBLIGACIONES"]
            netas_df = df[df["nv1"] == "RESERVAS NETAS"]
            if not brutas_df.empty and not netas_df.empty:
                brutas_tot = brutas_df[brutas_df["nv2"] == "TOTAL"].groupby("fecha")["valor"].first()
                oblig_tot = obligaciones_df.groupby("fecha")["valor"].sum() if not obligaciones_df.empty else pd.Series(0.0, index=brutas_tot.index)
                netas_tot = netas_df.groupby("fecha")["valor"].first()

                reconciled_netas = pd.concat([brutas_tot, oblig_tot, netas_tot], axis=1, keys=["brutas", "oblig", "netas"]).dropna()
                if not reconciled_netas.empty:
                    calc_netas = reconciled_netas["brutas"] - reconciled_netas["oblig"]
                    diff_netas = (reconciled_netas["netas"] - calc_netas).abs()
                    max_diff_netas = diff_netas.max()
                    if max_diff_netas > TOLERANCE:
                        print(f"Validation error in RESERVAS NETAS: max diff {max_diff_netas} exceeds {TOLERANCE}")
                        return True

            return False

        except Exception as error:
            print(f"Validation error: {error}")
            traceback.print_exc()
            return True


# Alias for compatibility with unit test runners expecting Robot
Robot = D_BO_000000463_01

Executor_D_BO_000000463_01 = D_BO_000000463_01
Robot = D_BO_000000463_01
