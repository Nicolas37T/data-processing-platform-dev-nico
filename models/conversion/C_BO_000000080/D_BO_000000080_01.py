import os
import re
import calendar
import traceback
import pandas as pd
from typing import Tuple

try:
    from models.conversion.Conversion_Base import Conversion_Base
except ImportError:
    class Conversion_Base:
        pass

try:
    from models.conversion.tools.conversion_tools import to_numeric_datax
except ImportError:
    try:
        from conversion_tools import to_numeric_datax
    except ImportError:
        pass


class D_BO_000000080_01(Conversion_Base):
    """
    DATAX Platform - Conversion Robot for D_BO_000000080_01
    Report: Cuentas Monetarias de No Bancos
    """

    def extraction(
        self,
        file_path: str,
        key_words: str = "",
        template_path: str = "",
        page_number: int = 1,
        format: str = "%Y-%m-%d"
    ) -> Tuple[dict, pd.DataFrame]:
        """
        Extracts monetary non-bank accounts report from Excel files into normalized melted format.
        """
        try:
            file_name = os.path.basename(file_path)

            def resolve_report_date(sample_text: str = "") -> str:
                # 1. Try to extract date from report text
                if sample_text:
                    month_map = {
                        "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
                        "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
                        "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12
                    }
                    pattern = r"(\d{1,2})\s+DE\s+([A-Za-z]+)\s+DE\s+(\d{4})"
                    match = re.search(pattern, sample_text, re.IGNORECASE)
                    if match:
                        day_str, month_str, year_str = match.groups()
                        month_num = month_map.get(month_str.lower())
                        if month_num:
                            return f"{int(year_str):04d}-{month_num:02d}-{int(day_str):02d}"

                # 2. Try to extract from file name (YYYYMM or YYYY-MM-DD)
                fn_match = re.search(r"(\d{4})[-_]?(\d{2})", file_name)
                if fn_match:
                    year = int(fn_match.group(1))
                    month = int(fn_match.group(2))
                    last_day = calendar.monthrange(year, month)[1]
                    return f"{year:04d}-{month:02d}-{last_day:02d}"

                # Fallback to current year-month last day
                today = pd.Timestamp.now()
                last_day = calendar.monthrange(today.year, today.month)[1]
                return f"{today.year:04d}-{today.month:02d}-{last_day:02d}"

            import openpyxl

            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]

            titles = []
            header_row_idx = None
            currencies = {}
            consolidado_title = "C O N S O L I D A D O"
            report_date = None
            exchange_rate_title = None
            denomination_title = None

            # Detect headers, titles, and currencies dynamically
            title_parts = []
            for row_idx in range(1, min(15, sheet.max_row + 1)):
                cell_a_val = sheet.cell(row_idx, 1).value
                cell_a_str = str(cell_a_val).strip() if cell_a_val is not None else ""

                if cell_a_str:
                    if "CUENTAS MONETARIAS" in cell_a_str.upper() or cell_a_str.upper() == "NO BANCOS":
                        title_parts.append(cell_a_str)
                    elif "CORRESPONDIENTE" in cell_a_str.upper():
                        if title_parts:
                            full_main_title = " ".join(title_parts).strip()
                            titles.append(full_main_title)
                            title_parts = []
                        titles.append(cell_a_str)
                        if report_date is None:
                            report_date = resolve_report_date(cell_a_str)
                    elif "TIPO DE CAMBIO" in cell_a_str.upper():
                        exchange_rate_title = cell_a_str

                # Locate CONSOLIDADO anchor and denomination subtitle
                for col_idx in range(1, 12):
                    val = sheet.cell(row_idx, col_idx).value
                    if val is not None:
                        val_str = str(val).strip()
                        if "MILLONES DE BOLIVIANOS" in val_str.upper():
                            denomination_title = val_str
                        if "CONSOLIDADO" in val_str.upper().replace(" ", ""):
                            consolidado_title = val_str

                # Locate currency column headers (MN, ME, MV, UFV, TOTAL)
                row_vals = [
                    (c, str(sheet.cell(row_idx, c).value).strip())
                    for c in range(1, 12)
                    if sheet.cell(row_idx, c).value is not None
                ]
                detected_curr = {
                    col: val.upper()
                    for col, val in row_vals
                    if val.upper() in {"MN", "ME", "MV", "UFV", "TOTAL"}
                }
                if len(detected_curr) >= 3:
                    header_row_idx = row_idx
                    currencies = detected_curr

            if title_parts and not any("CUENTAS MONETARIAS" in t for t in titles):
                titles.insert(0, " ".join(title_parts).strip())

            if exchange_rate_title and exchange_rate_title not in titles:
                titles.append(exchange_rate_title)
            if denomination_title and denomination_title not in titles:
                titles.append(denomination_title)

            if not report_date:
                report_date = resolve_report_date()

            if not header_row_idx or not currencies:
                return {"file_name": file_name, "titles": titles, "page_number": int(page_number)}, pd.DataFrame()

            records = []
            current_nv1 = None
            current_nv2 = None
            current_nv3 = None

            roman_pattern = re.compile(r"^\s*(I|II|III|IV|V|VI|VII|VIII|IX|X)\.?\s*", re.IGNORECASE)
            footer_tokens = ["FUENTE", "ELABORACI", "NOTAS"]

            # Parse hierarchical data rows
            for row_idx in range(header_row_idx + 1, sheet.max_row + 1):
                cell_a = sheet.cell(row_idx, 1).value
                cell_b = sheet.cell(row_idx, 2).value
                cell_c = sheet.cell(row_idx, 3).value

                str_a = str(cell_a).strip() if cell_a is not None else ""
                str_b = str(cell_b).strip() if cell_b is not None else ""
                str_c = str(cell_c).strip() if cell_c is not None else ""

                if not any([str_a, str_b, str_c]):
                    continue

                if any(token in str_a.upper() for token in footer_tokens):
                    break

                # Skip aggregate summary banner like TOTAL DEPÓSITOS
                if "TOTAL DEP" in str_b.upper():
                    continue

                # Determine hierarchy based on column placement and indentations
                if roman_pattern.match(str_a):
                    if str_b and re.match(r"^\s*(I|II|III|IV|V|VI|VII|VIII|IX|X)\.?\s*$", str_a, re.IGNORECASE):
                        current_nv1 = f"{str_a} {str_b}"
                    else:
                        current_nv1 = re.sub(r"[ \t]+", " ", str_a).strip()
                    current_nv2 = None
                    current_nv3 = None
                    nv1, nv2, nv3, nv4 = current_nv1, None, None, None
                elif str_b:
                    current_nv2 = str_b
                    current_nv3 = None
                    nv1, nv2, nv3, nv4 = current_nv1, current_nv2, None, None
                elif str_c:
                    cell_obj = sheet.cell(row_idx, 3)
                    indent_prop = cell_obj.alignment.indent if cell_obj.alignment else 0
                    raw_c_str = str(cell_c)
                    space_count = len(raw_c_str) - len(raw_c_str.lstrip(" "))

                    if indent_prop >= 1 or space_count >= 1:
                        nv1 = current_nv1
                        nv2 = current_nv2
                        nv3 = current_nv3
                        nv4 = str_c
                    else:
                        current_nv3 = str_c
                        nv1 = current_nv1
                        nv2 = current_nv2
                        nv3 = current_nv3
                        nv4 = None
                else:
                    continue

                raw_values = [sheet.cell(row_idx, col).value for col in currencies.keys()]
                has_numeric_data = any(
                    isinstance(v, (int, float)) and v is not None for v in raw_values
                )
                if not has_numeric_data:
                    continue

                for col_idx, currency_name in currencies.items():
                    val = sheet.cell(row_idx, col_idx).value
                    records.append({
                        "nv1": str(nv1) if nv1 is not None else None,
                        "nv2": str(nv2) if nv2 is not None else None,
                        "nv3": str(nv3) if nv3 is not None else None,
                        "nv4": str(nv4) if nv4 is not None else None,
                        "nv5": str(consolidado_title) if consolidado_title is not None else None,
                        "nv6": str(currency_name) if currency_name is not None else None,
                        "fecha": str(report_date),
                        "valor": val if val is not None else 0.0,
                    })

            if not records:
                return {"file_name": file_name, "titles": titles, "page_number": int(page_number)}, pd.DataFrame()

            df_melted = pd.DataFrame(records)

            # Ensure numeric conversion
            df_melted["valor"] = to_numeric_datax(df_melted["valor"], decimal_separator=".")

            # Ensure string or NaN compliance for metadata columns
            for col in ["nv1", "nv2", "nv3", "nv4", "nv5", "nv6", "fecha"]:
                df_melted[col] = df_melted[col].where(df_melted[col].notna(), None)

            metadata_dict = {
                "file_name": file_name,
                "titles": titles if titles else ["CUENTAS MONETARIAS DE NO BANCOS"],
                "page_number": int(page_number)
            }

            return metadata_dict, df_melted

        except Exception as error:
            print(f"Could not extract report from {file_path}: {error}")
            traceback.print_exc()
            return "", pd.DataFrame()

    def validate_data_results(
        self,
        dataframe: pd.DataFrame,
        decimal_separator: str = ",",
        TOLERANCE: float = 6.0
    ) -> bool:
        """
        Validates accounting reconciliation:
        1. Horizontal consistency: MN + ME + MV + UFV == TOTAL per row.
        2. Vertical consistency:
           a) Section roots equal the sum of their direct chapter categories.
           b) Specific sub-account components equal their respective parent total.

        Returns:
            False: Validation successful (no discrepancies detected).
            True: Discrepancies detected.
        """
        try:
            if not isinstance(dataframe, pd.DataFrame) or dataframe.empty:
                return True
            if "valor" not in dataframe.columns or "fecha" not in dataframe.columns:
                return True

            df = dataframe.copy()
            if pd.api.types.is_numeric_dtype(df["valor"]):
                df["valor_num"] = df["valor"].astype(float)
            else:
                df["valor_num"] = to_numeric_datax(df["valor"], decimal_separator=decimal_separator)

            if df["valor_num"].isna().all():
                return False

            # 1. Horizontal Reconciliation: MN + ME + MV + UFV == TOTAL
            group_keys = [c for c in ["nv1", "nv2", "nv3", "nv4", "fecha"] if c in df.columns]
            for keys, grp in df.groupby(group_keys, dropna=False):
                total_row = grp[grp["nv6"] == "TOTAL"]
                components = grp[grp["nv6"] != "TOTAL"]
                if total_row.empty or components.empty:
                    continue
                expected_total = total_row["valor_num"].iloc[0]
                calculated_total = components["valor_num"].sum()
                diff = abs(calculated_total - expected_total)
                if diff > TOLERANCE:
                    print(
                        f"Horizontal discrepancy at {keys}: "
                        f"Sum({calculated_total:.2f}) != TOTAL({expected_total:.2f}), Diff={diff:.2f}"
                    )
                    return True

            # 2. Vertical Reconciliation: Section root total == sum of chapter categories
            if "nv1" in df.columns and "nv2" in df.columns and "nv6" in df.columns:
                for (cut_date, currency, section_root), grp in df.groupby(["fecha", "nv6", "nv1"]):
                    root_row = grp[grp["nv2"].isna() & grp["nv3"].isna() & grp["nv4"].isna()]
                    chapters = grp[grp["nv2"].notna() & grp["nv3"].isna() & grp["nv4"].isna()]
                    if not root_row.empty and not chapters.empty:
                        root_val = root_row["valor_num"].iloc[0]
                        chapters_sum = chapters["valor_num"].sum()
                        diff = abs(root_val - chapters_sum)
                        if diff > TOLERANCE:
                            print(
                                f"Vertical discrepancy in section '{section_root}' ({currency}, {cut_date}): "
                                f"Root({root_val:.2f}) != Chapters({chapters_sum:.2f}), Diff={diff:.2f}"
                            )
                            return True

            # 3. Vertical Reconciliation: Sub-account Cartera Vigente/Vencida/Ejecución == Créditos
            if "nv3" in df.columns and "nv4" in df.columns:
                for (cut_date, currency), grp in df.groupby(["fecha", "nv6"]):
                    parent = grp[(grp["nv3"].str.strip() == "Créditos") & grp["nv4"].isna()]
                    children = grp[(grp["nv3"].str.strip() == "Créditos") & grp["nv4"].notna()]
                    if not parent.empty and not children.empty:
                        parent_val = parent["valor_num"].iloc[0]
                        children_sum = children["valor_num"].sum()
                        diff = abs(parent_val - children_sum)
                        if diff > TOLERANCE:
                            print(
                                f"Sub-account discrepancy in Créditos ({currency}, {cut_date}): "
                                f"Parent({parent_val:.2f}) != Children({children_sum:.2f}), Diff={diff:.2f}"
                            )
                            return True

            return False

        except Exception as error:
            print(f"Validation exception: {error}")
            traceback.print_exc()
            return True


Robot = D_BO_000000080_01
Executor_D_BO_000000080_01 = D_BO_000000080_01