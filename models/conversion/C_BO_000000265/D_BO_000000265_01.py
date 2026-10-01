from models.conversion.Conversion_Base import Conversion_Base
import os
import re
import calendar
import datetime
import traceback
from typing import Tuple
import pandas as pd
import openpyxl

from models.conversion.tools.conversion_tools import to_numeric_datax


class D_BO_000000265_01(Conversion_Base):
    """
    DATAX Conversion Robot for Report D_BO_000000265_01:
    Remesas de Trabajadores Recibidas según País de Origen.
    """

    def extraction(
        self,
        file_path: str,
        key_words: str = "",
        template_path: str = "",
        page_number: int = 1,
        format: str = "%Y-%m-%d"
    ) -> Tuple[dict, pd.DataFrame]:
        """
        Extracts workers' remittances by country of origin into the official DATAX DataFrame.

        Returns:
            Tuple[dict, pd.DataFrame]: (metadata_dict, df_melted)
        """
        try:
            file_name = os.path.basename(file_path)
            wb = openpyxl.load_workbook(file_path, data_only=True)

            month_map = {
                "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
                "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
                "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
                "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12
            }

            all_records = []
            clean_title = "REMESAS DE TRABAJADORES RECIBIDAS SEGÚN PAÍS DE ORIGEN"
            unit_title = None

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]

                # 1. Locate Titles and Headers Row dynamically
                header_row_idx = -1
                for r in range(1, min(20, sheet.max_row + 1)):
                    row_vals = [str(sheet.cell(r, c).value or "").strip() for c in range(1, min(25, sheet.max_column + 1))]
                    row_str = " ".join(v.upper() for v in row_vals)

                    if any("REMESAS DE TRABAJADORES RECIBIDAS" in v.upper() for v in row_vals):
                        for v in row_vals:
                            if "REMESAS DE TRABAJADORES RECIBIDAS" in v.upper():
                                clean_title = re.sub(r"\s*\d+$", "", v).strip()

                    for v in row_vals:
                        if "MILLONES DE" in v.upper():
                            unit_title = v.strip()

                    if "ALEMANIA" in row_str and any(k in row_str for k in ["REMESAS", "TOTAL"]):
                        header_row_idx = r
                        break

                if header_row_idx == -1:
                    continue

                # 2. Identify Column Boundaries
                start_col_idx = -1
                end_col_idx = -1
                date_col_idx = 2

                for c in range(1, sheet.max_column + 1):
                    val_str = str(sheet.cell(header_row_idx, c).value or "").strip().upper()
                    if "ALEMANIA" in val_str and start_col_idx == -1:
                        start_col_idx = c
                    if start_col_idx != -1 and any(k in val_str for k in ["REMESAS", "TOTAL"]):
                        end_col_idx = c

                if start_col_idx == -1 or end_col_idx == -1:
                    continue

                for c in range(1, start_col_idx):
                    val_str = str(sheet.cell(header_row_idx, c).value or "").strip().upper()
                    if any(k in val_str for k in ["MIGRAR", "MES", "FECHA", "AÑO", "ANO"]):
                        date_col_idx = c
                        break

                headers = []
                for c in range(start_col_idx, end_col_idx + 1):
                    h_val = str(sheet.cell(header_row_idx, c).value or "").strip()
                    if h_val.upper() == "TOTAL":
                        headers.append("Remesas")
                    else:
                        headers.append(h_val)

                # 3. Pre-scan for Latest Monthly Date (to anchor current-year preliminary tags)
                max_monthly_date = None
                for r in range(header_row_idx + 1, sheet.max_row + 1):
                    c_val = sheet.cell(r, date_col_idx).value
                    if c_val is None:
                        continue
                    if isinstance(c_val, (datetime.datetime, datetime.date)):
                        last_day = calendar.monthrange(c_val.year, c_val.month)[1]
                        d = datetime.date(c_val.year, c_val.month, last_day)
                        if max_monthly_date is None or d > max_monthly_date:
                            max_monthly_date = d
                    elif isinstance(c_val, str):
                        s = c_val.strip()
                        m1 = re.search(r"^(\d{4})-([A-Za-z]{3})", s)
                        if m1:
                            y = int(m1.group(1))
                            mo = month_map.get(m1.group(2).lower(), 1)
                            last_day = calendar.monthrange(y, mo)[1]
                            d = datetime.date(y, mo, last_day)
                            if max_monthly_date is None or d > max_monthly_date:
                                max_monthly_date = d
                        m2 = re.search(r"^([A-Za-z]{3})-(\d{2})", s)
                        if m2:
                            y_short = int(m2.group(2))
                            y = 2000 + y_short if y_short < 50 else 1900 + y_short
                            mo = month_map.get(m2.group(1).lower(), 1)
                            last_day = calendar.monthrange(y, mo)[1]
                            d = datetime.date(y, mo, last_day)
                            if max_monthly_date is None or d > max_monthly_date:
                                max_monthly_date = d

                # 4. Extract Data Rows
                for r in range(header_row_idx + 1, sheet.max_row + 1):
                    c_val = sheet.cell(r, date_col_idx).value
                    if c_val is None:
                        continue

                    parsed_date = None
                    nv2 = None

                    if isinstance(c_val, (datetime.datetime, datetime.date)):
                        last_day = calendar.monthrange(c_val.year, c_val.month)[1]
                        parsed_date = f"{c_val.year:04d}-{c_val.month:02d}-{last_day:02d}"
                        nv2 = None
                    else:
                        date_str = str(c_val).strip()
                        m_year = re.match(r"^(\d{4})\s*(p?)$", date_str, re.IGNORECASE)
                        if m_year:
                            year = int(m_year.group(1))
                            if max_monthly_date and year == max_monthly_date.year:
                                parsed_date = max_monthly_date.strftime("%Y-%m-%d")
                            else:
                                parsed_date = f"{year:04d}-12-31"
                            nv2 = date_str
                        else:
                            m_str = re.search(r"^(\d{4})-([A-Za-z]{3})\s*(p?)$", date_str, re.IGNORECASE)
                            if m_str:
                                year = int(m_str.group(1))
                                mo = month_map.get(m_str.group(2).lower(), 1)
                                last_day = calendar.monthrange(year, mo)[1]
                                parsed_date = f"{year:04d}-{mo:02d}-{last_day:02d}"
                                nv2 = None
                            else:
                                m_str_2 = re.search(r"^([A-Za-z]{3})-(\d{2})\s*(p?)$", date_str, re.IGNORECASE)
                                if m_str_2:
                                    y_short = int(m_str_2.group(2))
                                    year = 2000 + y_short if y_short < 50 else 1900 + y_short
                                    mo = month_map.get(m_str_2.group(1).lower(), 1)
                                    last_day = calendar.monthrange(year, mo)[1]
                                    parsed_date = f"{year:04d}-{mo:02d}-{last_day:02d}"
                                    nv2 = None

                    if parsed_date:
                        for idx, h in enumerate(headers):
                            col_idx = start_col_idx + idx
                            val = sheet.cell(r, col_idx).value
                            all_records.append({
                                "nv1": h,
                                "nv2": nv2,
                                "fecha": parsed_date,
                                "valor": val
                            })

            if not all_records:
                return {"file_name": file_name, "titles": [clean_title], "page_number": int(page_number)}, pd.DataFrame()

            df_melted = pd.DataFrame(all_records)

            # Ensure numeric conversion
            df_melted["valor"] = to_numeric_datax(df_melted["valor"], decimal_separator=",")
            df_melted = df_melted.dropna(subset=["valor"])

            # Deduplicate if multi-sheet
            df_melted = df_melted.drop_duplicates(subset=["nv1", "nv2", "fecha"])

            # Ensure strict columns ordering and string/None compliance
            for col in ["nv1", "nv2", "fecha"]:
                df_melted[col] = df_melted[col].where(df_melted[col].notna(), None)

            df_melted = df_melted[["nv1", "nv2", "fecha", "valor"]]

            titles = [clean_title]
            if unit_title and unit_title not in titles:
                titles.append(unit_title)

            metadata_dict = {
                "file_name": file_name,
                "titles": titles,
                "page_number": int(page_number) if page_number and int(page_number) > 0 else 1
            }

            return metadata_dict, df_melted

        except Exception as error:
            print(f"Error extracting report from {file_path}: {error}")
            traceback.print_exc()
            return "", pd.DataFrame()

    def validate_data_results(
        self,
        dataframe: pd.DataFrame,
        decimal_separator: str = ",",
        TOLERANCE: float = 6.0
    ) -> bool:
        """
        Validates mathematical completeness and accounting reconciliation:
        For each date (and nv2 category), the sum of all individual countries
        plus 'Otros' must equal the total 'Remesas' within TOLERANCE.

        Returns:
            False: Validation successful (data complete and reconciles).
            True: Discrepancy detected.
        """
        try:
            if not isinstance(dataframe, pd.DataFrame) or dataframe.empty:
                return True
            if "valor" not in dataframe.columns or "fecha" not in dataframe.columns:
                return True

            df = dataframe.copy()
            if pd.api.types.is_numeric_dtype(df["valor"]):
                df["valor_num"] = df["valor"].astype(float)
            else:
                df["valor_num"] = to_numeric_datax(df["valor"], decimal_separator=decimal_separator).fillna(0)

            group_cols = ["fecha", "nv2"] if "nv2" in df.columns else ["fecha"]
            for dt, group in df.groupby(group_cols, dropna=False):
                remesas_row = group[group["nv1"].str.upper().isin(["REMESAS", "TOTAL"])]
                if remesas_row.empty:
                    continue

                expected_total = remesas_row["valor_num"].iloc[0]
                calc_total = group[~group["nv1"].str.upper().isin(["REMESAS", "TOTAL"])]["valor_num"].sum()

                diff = abs(calc_total - expected_total)
                if diff > TOLERANCE:
                    print(
                        f"Validation discrepancy at date {dt}: "
                        f"Sum of countries ({calc_total:.4f}) != Remesas total ({expected_total:.4f}), Diff={diff:.4f}"
                    )
                    return True

            return False

        except Exception as error:
            print(f"Validation exception occurred: {error}")
            traceback.print_exc()
            return True


Robot = D_BO_000000265_01
Executor_D_BO_000000265_01 = D_BO_000000265_01