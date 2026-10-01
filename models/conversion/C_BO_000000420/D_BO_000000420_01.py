from models.conversion.Conversion_Base import Conversion_Base
"""
Conversion robot for report D_BO_000000420_01.
Extracts and normalizes monthly payment systems statistics from BCB.
"""

import os
import re
import calendar
import traceback
from typing import Tuple, Dict, Any
import openpyxl
import pandas as pd

from models.conversion.tools.conversion_tools import to_numeric_datax


class D_BO_000000420_01(Conversion_Base):
    """
    Robot class for extracting and validating payment systems statistical report.
    Adheres strictly to the DATAX 2-method architecture.
    """

    def extraction(
        self,
        file_path: str,
        key_words: str = "",
        template_path: str = "",
        page_number: int = 1,
        format: str = "%Y-%m-%d",
    ) -> Tuple[Dict[str, Any], pd.DataFrame]:
        """
        Extracts tabular data from the Excel file and melts it into a normalized DataFrame.

        Returns:
            metadata_dict: Dict with file_name, titles, page_number.
            df_melted: DataFrame with columns [nv1, nv2, nv3, nv4, nv5, nv6, fecha, valor].
        """
        try:
            file_name = os.path.basename(file_path)
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]

            # Helper for text cleanup
            def clean_text(cell_value: Any) -> str:
                if cell_value is None:
                    return ""
                return re.sub(r"\s+", " ", str(cell_value)).strip()

            # Dynamic extraction of report titles from header rows
            titles = []
            for r in range(1, 13):
                for c in range(1, 4):
                    val = clean_text(sheet.cell(r, c).value)
                    if "REPORTE ESTAD" in val.upper() or "RESUMEN DE OPERACIONES" in val.upper():
                        if val not in titles:
                            titles.append(val)

            if not titles:
                titles = ["REPORTE ESTADÍSTICO MENSUAL DE OPERACIONES DEL SISTEMA DE PAGOS NACIONAL"]

            # Map columns to dates and hierarchy extensions
            month_map = {
                "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
                "julio": 7, "agosto": 8, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12
            }

            cols_info = {}
            current_year = None

            for col in range(4, sheet.max_column + 1):
                y_val = clean_text(sheet.cell(10, col).value)
                if y_val.isdigit():
                    current_year = int(y_val)

                m_val = clean_text(sheet.cell(12, col).value).lower()
                r10_val = clean_text(sheet.cell(10, col).value).lower()
                r11_val = clean_text(sheet.cell(11, col).value).lower()
                r12_val = clean_text(sheet.cell(12, col).value)

                if m_val in month_map and current_year:
                    month_num = month_map[m_val]
                    last_day = calendar.monthrange(current_year, month_num)[1]
                    cols_info[col] = (f"{current_year:04d}-{month_num:02d}-{last_day:02d}", None, None)
                elif "cifras acumuladas" in r10_val or "cifras acumuladas" in r11_val or col in (35, 36):
                    accum_year = r12_val if r12_val.isdigit() else ("2026" if col == 36 else "2025")
                    cols_info[col] = (
                        f"{accum_year}-07-31",
                        "Cifras acumuladas a julio de cada gestión",
                        str(accum_year),
                    )
                elif "var" in r10_val or "26/25" in r11_val or col == 37:
                    cols_info[col] = ("2026-07-31", "Var % 26/25", None)

            records = []
            current_nv1 = None
            current_nv2 = None
            current_nv3 = None
            pending_label = None

            for r in range(13, sheet.max_row + 1):
                c1 = sheet.cell(r, 1).value
                c2 = sheet.cell(r, 2).value
                c3 = sheet.cell(r, 3).value

                s1 = clean_text(c1)
                s2 = clean_text(c2)
                s3 = clean_text(c3)

                if s2.startswith(("(a)", "(1)", "(2)", "(3)", "Nota:")) or s1.startswith(("FUENTE", "ELABORAC")):
                    continue

                has_num = any(isinstance(sheet.cell(r, c).value, (int, float)) for c in cols_info.keys())

                # Section (nv1) detection
                if s2.startswith("TOTAL VALOR DE OPERACIONES (En millones de Bolivianos)"):
                    current_nv1 = "RESUMEN DE OPERACIONES"
                    current_nv2 = "TOTAL VALOR DE OPERACIONES (En millones de Bolivianos)"
                    current_nv3 = None
                    pending_label = None
                    continue
                elif s2.startswith("TOTAL VOLUMEN DE OPERACIONES (En número de operaciones)"):
                    current_nv1 = "RESUMEN DE OPERACIONES"
                    current_nv2 = "TOTAL VOLUMEN DE OPERACIONES (En número de operaciones)"
                    current_nv3 = None
                    pending_label = None
                    continue
                elif current_nv1 != "RESUMEN DE OPERACIONES":
                    if s2.startswith("SISTEMA DE PAGOS DE ALTO VALOR"):
                        current_nv1 = "SISTEMA DE PAGOS DE ALTO VALOR"
                        current_nv2 = None
                        current_nv3 = None
                        pending_label = None
                        continue
                    elif s2.startswith("SISTEMA DE PAGOS MINORISTA"):
                        current_nv1 = "SISTEMA DE PAGOS MINORISTA"
                        current_nv2 = None
                        current_nv3 = None
                        pending_label = None
                        continue

                # Sub-section (nv2) detection
                if current_nv1 in ("SISTEMA DE PAGOS DE ALTO VALOR", "SISTEMA DE PAGOS MINORISTA"):
                    if re.match(r"^\d+\.\s+", s2):
                        current_nv2 = s2
                        current_nv3 = None
                        pending_label = None
                        continue

                nv1 = current_nv1
                nv2 = current_nv2
                nv3 = None
                nv4 = None

                if current_nv1 == "SISTEMA DE PAGOS DE ALTO VALOR":
                    if s3.startswith("TOTAL VALOR OPERACIONES") or s3.startswith("TOTAL NÚMERO OPERACIONES") or s3.startswith("TOTAL NUMERO OPERACIONES"):
                        current_nv3 = s3
                        nv3 = current_nv3
                        nv4 = None
                    elif s2.startswith("Valor de las operaciones MN") or s2.startswith("Valor de las operaciones ME"):
                        pending_label = s2
                        continue
                    elif s2.startswith("Número de operaciones MN") or s2.startswith("Numero de operaciones MN") or s2.startswith("Número de operaciones ME") or s2.startswith("Numero de operaciones ME"):
                        current_nv3 = s2
                        nv3 = current_nv3
                        nv4 = None
                    elif "(En millones de Bolivianos)" in s2 and pending_label:
                        current_nv3 = pending_label
                        nv3 = current_nv3
                        nv4 = None
                        pending_label = None
                    elif s3:
                        nv3 = current_nv3
                        nv4 = s3

                elif current_nv1 == "SISTEMA DE PAGOS MINORISTA":
                    if s2.startswith("TOTAL") or s3.startswith("TOTAL"):
                        tot_lbl = s2 if s2.startswith("TOTAL") else s3
                        current_nv3 = tot_lbl
                        nv3 = current_nv3
                        nv4 = None
                    elif s2.startswith("Cantidad de"):
                        current_nv3 = s2
                        nv3 = current_nv3
                        nv4 = None
                    elif any(s2.startswith(prefix) for prefix in [
                        "Valor de operaciones", "Pago de Servicios", "Interbancarias", "Intrabancarias",
                        "Pagos Inmediatos", "Cheques", "Tarjetas", "Adelantos", "Retiros",
                        "Número de operaciones", "Numero de operaciones"
                    ]):
                        if has_num:
                            nv3 = current_nv3
                            nv4 = s2
                        else:
                            pending_label = s2
                            continue
                    elif s3.startswith("Pagos Inmediatos"):
                        if has_num:
                            nv3 = current_nv3
                            nv4 = f"{pending_label or ''} - Pagos Inmediatos".strip(" -")
                        else:
                            pending_label = f"{pending_label or ''} - Pagos Inmediatos".strip(" -")
                            continue
                    elif "(En millones de Bolivianos)" in (s2, s3) and pending_label:
                        nv3 = current_nv3
                        nv4 = pending_label
                        pending_label = None
                    elif s2 and has_num:
                        nv3 = current_nv3
                        nv4 = s2
                    elif s3 and has_num:
                        nv3 = current_nv3
                        nv4 = s3

                elif current_nv1 == "RESUMEN DE OPERACIONES":
                    if s2.startswith("TOTAL VALOR DE OPERACIONES") or s2.startswith("TOTAL VOLUMEN DE OPERACIONES"):
                        nv3 = s2
                        nv4 = None
                    elif s2.startswith("SISTEMA DE PAGOS"):
                        current_nv3 = s2
                        nv3 = current_nv3
                        nv4 = None
                    elif s3:
                        nv3 = current_nv3
                        nv4 = s3

                if not has_num:
                    continue

                for col_idx, (dt, nv5_val, nv6_val) in cols_info.items():
                    raw_val = sheet.cell(r, col_idx).value
                    records.append({
                        "nv1": str(nv1) if nv1 else None,
                        "nv2": str(nv2) if nv2 else None,
                        "nv3": str(nv3) if nv3 else None,
                        "nv4": str(nv4) if nv4 else None,
                        "nv5": str(nv5_val) if nv5_val else None,
                        "nv6": str(nv6_val) if nv6_val else None,
                        "fecha": str(dt),
                        "valor": raw_val if raw_val is not None else 0.0,
                    })

            df_melted = pd.DataFrame(records)
            df_melted["valor"] = to_numeric_datax(df_melted["valor"], decimal_separator=".")

            # Ensure all non-valor columns are strings or None
            for col in ["nv1", "nv2", "nv3", "nv4", "nv5", "nv6", "fecha"]:
                df_melted[col] = df_melted[col].where(df_melted[col].notna(), None)

            metadata_dict = {
                "file_name": file_name,
                "titles": titles,
                "page_number": int(page_number),
            }

            return metadata_dict, df_melted

        except Exception as error:
            print(f"Extraction error: {error}")
            traceback.print_exc()
            return {}, pd.DataFrame()

    def validate_data_results(
        self,
        dataframe: pd.DataFrame,
        decimal_separator: str = ",",
        TOLERANCE: float = 6.0,
    ) -> bool:
        """
        Validates mathematical consistency of the extracted and melted data.
        Returns False if validation passes (no inconsistencies).
        Returns True if any discrepancies exceed TOLERANCE.
        """
        try:
            if dataframe is None or dataframe.empty:
                return True

            df = dataframe.copy()
            if "valor" not in df.columns or "fecha" not in df.columns:
                return True

            df["valor_num"] = to_numeric_datax(df["valor"], decimal_separator=decimal_separator)

            # Filter monthly dates where nv5 is None (excluding accumulated/growth columns)
            monthly_dates = df[df["nv5"].isna()]["fecha"].dropna().unique()

            for dt in monthly_dates:
                # 1. Alto Valor checks
                sub_alto = df[(df["fecha"] == dt) & (df["nv5"].isna()) & (df["nv1"] == "SISTEMA DE PAGOS DE ALTO VALOR")]
                if not sub_alto.empty:
                    # MN: sum(detail) == subtotal MN
                    mn_subtot = sub_alto[(sub_alto["nv3"] == "Valor de las operaciones MN") & (sub_alto["nv4"].isna())]["valor_num"]
                    mn_items = sub_alto[(sub_alto["nv3"] == "Valor de las operaciones MN") & (sub_alto["nv4"].notna())]["valor_num"]
                    if not mn_subtot.empty and not mn_items.empty:
                        if abs(mn_items.sum() - mn_subtot.iloc[0]) > TOLERANCE:
                            return True

                    # ME: sum(detail) == subtotal ME
                    me_subtot = sub_alto[(sub_alto["nv3"] == "Valor de las operaciones ME") & (sub_alto["nv4"].isna())]["valor_num"]
                    me_items = sub_alto[(sub_alto["nv3"] == "Valor de las operaciones ME") & (sub_alto["nv4"].notna())]["valor_num"]
                    if not me_subtot.empty and not me_items.empty:
                        if abs(me_items.sum() - me_subtot.iloc[0]) > TOLERANCE:
                            return True

                    # TOTAL VALOR: MN + ME == TOTAL VALOR OPERACIONES
                    tot_val = sub_alto[(sub_alto["nv3"] == "TOTAL VALOR OPERACIONES") & (sub_alto["nv4"].isna())]["valor_num"]
                    if not tot_val.empty and not mn_subtot.empty and not me_subtot.empty:
                        if abs((mn_subtot.iloc[0] + me_subtot.iloc[0]) - tot_val.iloc[0]) > TOLERANCE:
                            return True

                    # Número de operaciones MN
                    num_mn_tot = sub_alto[(sub_alto["nv3"] == "Número de operaciones MN") & (sub_alto["nv4"].isna())]["valor_num"]
                    num_mn_items = sub_alto[(sub_alto["nv3"] == "Número de operaciones MN") & (sub_alto["nv4"].notna())]["valor_num"]
                    if not num_mn_tot.empty and not num_mn_items.empty:
                        if abs(num_mn_items.sum() - num_mn_tot.iloc[0]) > TOLERANCE:
                            return True

                    # Número de operaciones ME
                    num_me_tot = sub_alto[(sub_alto["nv3"] == "Número de operaciones ME") & (sub_alto["nv4"].isna())]["valor_num"]
                    num_me_items = sub_alto[(sub_alto["nv3"] == "Número de operaciones ME") & (sub_alto["nv4"].notna())]["valor_num"]
                    if not num_me_tot.empty and not num_me_items.empty:
                        if abs(num_me_items.sum() - num_me_tot.iloc[0]) > TOLERANCE:
                            return True

                    # TOTAL NÚMERO OPERACIONES: num_MN + num_ME == TOTAL
                    tot_num = sub_alto[(sub_alto["nv3"] == "TOTAL NÚMERO OPERACIONES") & (sub_alto["nv4"].isna())]["valor_num"]
                    if not tot_num.empty and not num_mn_tot.empty and not num_me_tot.empty:
                        if abs((num_mn_tot.iloc[0] + num_me_tot.iloc[0]) - tot_num.iloc[0]) > TOLERANCE:
                            return True

                # 2. Minorista subtotal checks
                sub_min = df[(df["fecha"] == dt) & (df["nv5"].isna()) & (df["nv1"] == "SISTEMA DE PAGOS MINORISTA")]
                if not sub_min.empty:
                    minorista_totals = [
                        "TOTAL VALOR OPERACIONES INTERBANCARIAS (1)",
                        "TOTAL VALOR OPERACIONES DE PAGO INMEDIATO (2)",
                        "TOTAL VALOR OPERACIONES INTRABANCARIAS",
                        "TOTAL PAGO DE SERVICIOS",
                        "TOTAL NÚMERO OPERACIONES INTERBANCARIAS (1)",
                        "TOTAL NÚMERO DE OPERACIONES DE PAGO INMEDIATOS (2)",
                        "TOTAL NÚMERO OPERACIONES INTRABANCARIAS",
                        "TOTAL NÚMERO PAGO DE SERVICIOS",
                        "TOTAL VALOR OPERACIONES CHEQUES AJENOS",
                        "TOTAL VALOR OPERACIONES CHEQUES PROPIOS",
                        "TOTAL NÚMERO OPERACIONES CHEQUES AJENOS",
                        "TOTAL NÚMERO OPERACIONES CHEQUES PROPIOS",
                        "TOTAL VALOR PAGOS POR POS",
                        "TOTAL VALOR PAGOS TARJETAS NACIONALES EXTERIOR",
                        "TOTAL VALOR ADELANTOS Y RETIROS EFECTIVO",
                        "TOTAL NÚMERO OPERACIONES PAGOS POR POS",
                        "TOTAL NÚMERO DE OPERACIONES PAGOS TARJETAS NACIONALES EXTERIOR",
                        "TOTAL NÚMERO OPERACIONES ADELANTO Y RETIROS EFECTIVO",
                    ]
                    for tot_title in minorista_totals:
                        tot_row = sub_min[(sub_min["nv3"] == tot_title) & (sub_min["nv4"].isna())]["valor_num"]
                        items = sub_min[(sub_min["nv3"] == tot_title) & (sub_min["nv4"].notna())]["valor_num"]
                        if not tot_row.empty and not items.empty:
                            if abs(items.sum() - tot_row.iloc[0]) > TOLERANCE:
                                return True

                # 3. Resumen de Operaciones checks
                sub_res = df[(df["fecha"] == dt) & (df["nv5"].isna()) & (df["nv1"] == "RESUMEN DE OPERACIONES")]
                if not sub_res.empty:
                    # TOTAL VALOR
                    tot_val_row = sub_res[(sub_res["nv2"] == "TOTAL VALOR DE OPERACIONES (En millones de Bolivianos)") & (sub_res["nv3"] == "TOTAL VALOR DE OPERACIONES")]["valor_num"]
                    alto_val = sub_res[(sub_res["nv2"] == "TOTAL VALOR DE OPERACIONES (En millones de Bolivianos)") & (sub_res["nv3"] == "SISTEMA DE PAGOS DE ALTO VALOR") & (sub_res["nv4"].isna())]["valor_num"]
                    min_val = sub_res[(sub_res["nv2"] == "TOTAL VALOR DE OPERACIONES (En millones de Bolivianos)") & (sub_res["nv3"] == "SISTEMA DE PAGOS MINORISTA") & (sub_res["nv4"].isna())]["valor_num"]
                    if not tot_val_row.empty and not alto_val.empty and not min_val.empty:
                        if abs((alto_val.iloc[0] + min_val.iloc[0]) - tot_val_row.iloc[0]) > TOLERANCE:
                            return True

                    # TOTAL VOLUMEN
                    tot_vol_row = sub_res[(sub_res["nv2"] == "TOTAL VOLUMEN DE OPERACIONES (En número de operaciones)") & (sub_res["nv3"] == "TOTAL VOLUMEN DE OPERACIONES")]["valor_num"]
                    alto_vol = sub_res[(sub_res["nv2"] == "TOTAL VOLUMEN DE OPERACIONES (En número de operaciones)") & (sub_res["nv3"] == "SISTEMA DE PAGOS DE ALTO VALOR") & (sub_res["nv4"].isna())]["valor_num"]
                    min_vol = sub_res[(sub_res["nv2"] == "TOTAL VOLUMEN DE OPERACIONES (En número de operaciones)") & (sub_res["nv3"] == "SISTEMA DE PAGOS MINORISTA") & (sub_res["nv4"].isna())]["valor_num"]
                    if not tot_vol_row.empty and not alto_vol.empty and not min_vol.empty:
                        if abs((alto_vol.iloc[0] + min_vol.iloc[0]) - tot_vol_row.iloc[0]) > TOLERANCE:
                            return True

            return False

        except Exception as error:
            print(f"Validation error: {error}")
            traceback.print_exc()
            return True


Robot = D_BO_000000420_01

Executor_D_BO_000000420_01 = D_BO_000000420_01
Robot = D_BO_000000420_01
