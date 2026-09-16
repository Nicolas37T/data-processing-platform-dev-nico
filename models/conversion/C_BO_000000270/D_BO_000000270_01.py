import os
import re
import calendar
import traceback
from typing import Tuple
import openpyxl
import pandas as pd

from models.conversion.Conversion_Base import Conversion_Base
from models.conversion.tools.conversion_tools import to_numeric_datax


class D_BO_000000270_01(Conversion_Base):
    """
    Conversion robot for report D_BO_000000270_01:
    'Saldo de la Deuda Pública Interna del TGN (En Bs.)'
    Source: Ministerio de Economía y Finanzas Públicas (MEFP) / BCB
    Target sheet: 'Saldo'
    """

    def extraction(
        self,
        file_path: str,
        key_words: str = "",
        template_path: str = "",
        page_number: int = 1,
        format: str = "%Y-%m-%d",
    ) -> Tuple[dict, pd.DataFrame]:
        try:
            file_name = os.path.basename(file_path)

            cut_year, cut_month = 2026, 12
            date_match = re.search(r"(\d{4})-(\d{2})-(\d{2})", file_name)
            if date_match:
                cut_year = int(date_match.group(1))
                cut_month = int(date_match.group(2))

            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet_name = "Saldo"
            if sheet_name not in wb.sheetnames:
                raise ValueError(
                    f"Sheet '{sheet_name}' not found in file '{file_name}'. Available: {wb.sheetnames}"
                )
            ws = wb[sheet_name]

            sheet_title = (
                ws.cell(2, 1).value
                or "TGN Deuda Pública Interna: Saldo, según instrumento"
            )
            year_cell = str(ws.cell(3, 1).value or "")
            year_match = re.search(r"(\d{4})", year_cell)
            if year_match and not date_match:
                cut_year = int(year_match.group(1))

            month_map = {
                "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
                "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12,
            }

            col_dates = {}
            header_row = 6
            for col in range(2, 15):
                val = ws.cell(header_row, col).value
                if not val:
                    continue
                h_str = str(val).strip().lower()

                m_yr = re.search(r"([a-z]{3})[-_ ]?(\d{2,4})", h_str)
                if m_yr:
                    m_str = m_yr.group(1)
                    yr = int(m_yr.group(2))
                    if yr < 100:
                        yr += 2000
                    m_num = month_map.get(m_str)
                else:
                    m_num = month_map.get(h_str[:3])
                    yr = cut_year

                if m_num:
                    if yr < cut_year or (yr == cut_year and m_num <= cut_month):
                        last_day = calendar.monthrange(yr, m_num)[1]
                        col_dates[col] = f"{yr:04d}-{m_num:02d}-{last_day:02d}"

            records = []

            def append_row_records(row_idx, nv1, nv2, nv3, nv4, nv5):
                for col, dt in col_dates.items():
                    raw_val = ws.cell(row_idx, col).value
                    records.append({
                        "nv1": str(nv1) if nv1 is not None else None,
                        "nv2": str(nv2) if nv2 is not None else None,
                        "nv3": str(nv3) if nv3 is not None else None,
                        "nv4": str(nv4) if nv4 is not None else None,
                        "nv5": str(nv5) if nv5 is not None else None,
                        "fecha": dt,
                        "valor": raw_val,
                    })

            top_title = "Deuda Pública Interna Total del TGN"
            curr_sector = None
            curr_inst = None

            for r in range(7, ws.max_row + 1):
                cell = ws.cell(r, 1)
                val = cell.value
                if not val:
                    continue

                text = str(val).strip()
                if text.lower().startswith(
                    ("(p)", "nota", "fuente", "elaboraci", "cierre presupuestario")
                ):
                    break

                bold = cell.font.bold if cell.font else False
                indent = cell.alignment.indent if cell.alignment else 0
                leading_spaces = len(str(val)) - len(str(val).lstrip(" "))

                if bold and indent == 0 and leading_spaces == 0:
                    if "total" in text.lower():
                        append_row_records(r, text, None, None, None, None)
                    else:
                        curr_sector = text
                        curr_inst = None
                    continue

                if bold and indent >= 1:
                    curr_inst = text
                    continue

                if not bold and curr_sector:
                    nv1 = top_title
                    nv2 = curr_sector
                    nv3 = curr_inst
                    nv4 = None
                    nv5 = None

                    if curr_sector == "Sector Público Financiero":
                        if curr_inst == "BCB":
                            if "cred.emerg" in text.lower():
                                nv4 = "BTs-Cred.Emerg."
                            elif "bts-neg" in text.lower() and "no" not in text.lower():
                                nv4 = "Bts-Neg."
                            elif "bts-no neg" in text.lower() or "no neg" in text.lower():
                                nv4 = "Bts-Neg."
                                nv5 = "spf_bts-no neg."
                            elif "hist" in text.lower() and '"a"' in text.lower():
                                nv4 = 'Deuda Hist-LT "A"'
                            elif "hist" in text.lower() and '"b"' in text.lower():
                                nv4 = 'Deuda Hist-LT "B"'
                            elif "títulos" in text.lower() or "titulos" in text.lower():
                                nv4 = "Títulos-BCB"
                        elif curr_inst == "Fondos":
                            if "no neg" in text.lower():
                                nv4 = "BTs-Fdos. No Neg."
                            else:
                                nv4 = "BTs-Fdos. Neg."

                    elif curr_sector == "Sector Público no Financiero":
                        if curr_inst == "Otros Públicos":
                            nv4 = "BTs-No Neg."

                    elif curr_sector == "Sector Privado":
                        if curr_inst == "AFPs":
                            nv4 = "Bonos-AFPs"
                        elif curr_inst and "mercado financiero" in curr_inst.lower():
                            nv3 = "AFPs"
                            nv4 = "MCDO. FINAN. (Subasta)"
                            if "bonos" in text.lower():
                                nv5 = 'Bonos "C"'
                            elif "letras" in text.lower():
                                nv5 = 'Letras "C"'
                        elif curr_inst == "Otros privados":
                            nv4 = "Bonos Privados"
                        elif curr_inst == "Fondos":
                            if "no neg" in text.lower():
                                nv4 = "sp-fdos-bts-no neg."
                            else:
                                nv4 = "sp-fdos-bts-neg."
                        elif curr_inst == "Tesoro directo":
                            nv4 = "BTs-Extrabursatil"

                    if nv1 and nv4:
                        append_row_records(r, nv1, nv2, nv3, nv4, nv5)

            df_data = pd.DataFrame(records)

            # Clean numeric column using to_numeric_datax
            df_data["valor"] = to_numeric_datax(df_data["valor"], decimal_separator=",").round()

            metadata_dict = {
                "file_name": file_name,
                "titles": [str(sheet_title).strip()],
                "page_number": int(page_number),
            }

            return metadata_dict, df_data

        except Exception as error:
            print(f"Could not extract report from {file_path}: {error}")
            traceback.print_exc()
            return ""

    def validate_data_results(
        self,
        dataframe: pd.DataFrame,
        decimal_separator: str = ",",
        TOLERANCE: float = 6.0,
    ) -> bool:
        """
        Valida que el aplanamiento de datos sea matematicamente consistente.
        Reconcilia que la suma de los instrumentos individuales (amarillos)
        sea igual al Total Deuda Publica Interna del TGN (naranja) para cada
        fecha de corte mensual dentro de la tolerancia admitida.

        Retorna:
            False: Si la validacion es exitosa (sin inconsistencias).
            True: Si se detecta un error de cuadre o discrepancia numerica.
        """
        try:
            if not isinstance(dataframe, pd.DataFrame) or dataframe.empty:
                return True
            if "valor" not in dataframe.columns or "fecha" not in dataframe.columns:
                return True

            df = dataframe.copy()
            df["valor_num"] = to_numeric_datax(df["valor"], decimal_separator)

            for dt, group in df.groupby("fecha"):
                total_mask = group["nv2"].isna()
                if not total_mask.any():
                    print(f"Validation error: Fila Total TGN no encontrada para fecha {dt}")
                    return True

                expected_total = group.loc[total_mask, "valor_num"].iloc[0]
                leaves = group.loc[~total_mask]
                calculated_sum = leaves["valor_num"].sum()

                diff = abs(calculated_sum - expected_total)
                if diff > TOLERANCE:
                    print(
                        f"Validation error en fecha {dt}: Suma de instrumentos ({calculated_sum:.2f}) "
                        f"difiere del Total TGN ({expected_total:.2f}) por {diff:.2f} > TOLERANCE ({TOLERANCE})"
                    )
                    return True

            return False
        except Exception as error:
            print(f"Validation exception: {error}")
            return True


Robot = D_BO_000000270_01